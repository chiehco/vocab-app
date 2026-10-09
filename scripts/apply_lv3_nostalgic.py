"""Stage the approved nostalgic entry, then commit validated App outputs.

The source workbook is read-only during staging. --publish-master requires a
separate passing scope audit and creates an exclusive backup before replacement.
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import openpyxl

from apply_lv3_approved_review import (
    CONTENT_FILES, NS, append_row, published_content_hash, read, replace_outputs, sha,
)
from build_sa_pack import build_sa_pack
from import_data import parse_examples, parse_senses, parse_words, write_json

BASE_MASTER_HASH = '0af6602355b3ef31cbb54cbd43fe412566c7582326e47684d2f71caa22af4bdb'
WORD_ID = 'W006086'
SENSE_ID = WORD_ID + '-1'
EXAMPLE_ID = 'EX-LV3U01-nostalgic'
WORD = 'nostalgic'
MEANING = '懷念往日時光的'
SENTENCE = 'Finding her old school notebook made her feel nostalgic.'
SENTENCE_ZH = '找到以前的學校筆記本，讓她懷念起往日時光。'
WORK_EXAMPLE_ID = 'ORIG-20260912-LV3U01-SUP-nostalgic'
APPROVAL_ID = 'Sentinel_2ff027612884819198ecf741c1308c75'
SOURCE_NOTE = '常春藤 LV3 Unit1 p.6 補充詞；2026-10-09 使用者批准 nostalgic／懷念往日時光的；沿用原工作例句，不改為 homesick。'
UNCHANGED_DATA = [name for name in CONTENT_FILES if name not in {'words.json', 'senses.json', 'examples.json'}]
OUTPUT_FILES = [
    'public/data/v1/words.json', 'public/data/v1/senses.json', 'public/data/v1/examples.json',
    'public/data/v1/meta.json', 'public/data/v1/sa-pack.json',
    'src/features/direct/curriculum.json', 'src/features/direct/wordCatalog.json',
    'src/features/vocabulary/lv3NostalgicReview.json',
]


def stage_master(source, staged):
    assert sha(source) == BASE_MASTER_HASH, 'Master changed; re-audit before applying'
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=False)
    try:
        ids = set()
        for sheet in workbook:
            for row in sheet.values:
                for value in row:
                    if isinstance(value, str):
                        ids.update(re.findall(r'\bW\d{6}\b', value))
        assert WORD_ID not in ids and max(ids) == 'W006085', 'Word ID conflict'
        assert len(workbook.sheetnames) == 64
        assert not any(WORD in (record.get('wordVariants') or []) for record in parse_words(workbook['input_words_單字主表'])), 'Already applied; do not duplicate'
        values = {
            'input_words_單字主表': [WORD_ID, WORD, 'LV3', 'adj.', MEANING, None, None, None, None, None, WORD, False, SOURCE_NOTE, 'reviewed', None, None],
            'input_senses_義項表': [SENSE_ID, WORD_ID, WORD, 'adj.', MEANING, None, None, WORD, '僅本次核准的懷念往日時光義；不等同 homesick。', 'reviewed'],
            'input_examples_例句表': [EXAMPLE_ID, WORD, 'adj.', MEANING, 'daily', SENTENCE, SENTENCE_ZH,
                'Finding her old school notebook made her feel _____.', WORD, 'LV3', 'reviewed',
                f'沿用 {WORK_EXAMPLE_ID}；原例句 provenance=authored_in_this_task（既有 2026-09-12 工作稿），非本次新撰。'],
        }
        counts = {name: len(list(workbook[name].values)) for name in values}
    finally:
        workbook.close()
    with ZipFile(source) as original:
        rels = {row.attrib['Id']: row.attrib['Target'] for row in ET.fromstring(original.read('xl/_rels/workbook.xml.rels'))}
        sheets = ET.fromstring(original.read('xl/workbook.xml')).find('s:sheets', NS)
        paths = {sheet.attrib['name']: rels[sheet.attrib[f'{{{NS["r"]}}}id']] for sheet in sheets}
        replacements, rows = {}, {}
        for name, row in values.items():
            path = paths[name]
            path = path.lstrip('/') if path.startswith('/') else 'xl/' + path
            replacements[path], rows[name] = append_row(original.read(path), row, counts[name])
        with ZipFile(staged, 'w') as result:
            for member in original.infolist():
                result.writestr(member, replacements.get(member.filename, original.read(member)))
    with ZipFile(source) as original, ZipFile(staged) as result:
        assert original.namelist() == result.namelist()
        changed = [name for name in original.namelist() if original.read(name) != result.read(name)]
        assert set(changed) == set(replacements)
    return dict(sourceSha256=BASE_MASTER_HASH, stagedSha256=sha(staged), rows=rows,
                changedWorkbookParts=changed, preservedSheets=64, priorCellsAndFormulas='byte-preserved')


def build_outputs(repo, master, out):
    out.mkdir(parents=True)
    staged = out / master.name
    evidence = stage_master(master, staged)
    data = repo / 'public/data/v1'
    workbook = openpyxl.load_workbook(staged, read_only=True, data_only=True)
    try:
        additions = [
            ('words.json', next(row for row in parse_words(workbook['input_words_單字主表']) if row['wordId'] == WORD_ID), 'wordId'),
            ('senses.json', next(row for row in parse_senses(workbook['input_senses_義項表']) if row['senseId'] == SENSE_ID), 'senseId'),
            ('examples.json', next(row for row in parse_examples(workbook['input_examples_例句表']) if row['exampleId'] == EXAMPLE_ID), 'exampleId'),
        ]
    finally:
        workbook.close()
    for filename, addition, key in additions:
        records = read(data / filename)
        assert not any(row[key] == addition[key] or (filename == 'words.json' and WORD in (row.get('wordVariants') or [])) for row in records), 'Already applied; do not duplicate'
        records.append(addition)
        records.sort(key=lambda row: row[key])
        write_json(data / filename, records)
    unit_path = repo / 'src/features/direct/curriculum.json'
    unit = read(unit_path)
    items = [item for item in unit['learningItems'] if item.get('displayWord') == WORD]
    assert len(items) == 1
    item = items[0]
    assert item['learningItemId'] == 'LI-LV3U01-nostalgic' and item['sensePos'] == 'adj.' and item['targetMeaningZh'] == MEANING
    example = item['originalExample']
    assert (example['workExampleId'], example['sentenceEn'], example['sentenceZh']) == (WORK_EXAMPLE_ID, SENTENCE, SENTENCE_ZH)
    assert item['officialWordId'] is None and item['officialSenseId'] is None and item['lexemeRef'] == 'LC-LV3U01-nostalgic'
    original_chapter_meaning = item.get('chapterMeaningsZh')
    assert original_chapter_meaning == '鄉愁的'
    item.update(lexemeRef=WORD_ID, officialWordId=WORD_ID, officialSenseId=SENSE_ID, chapterMeaningsZh=MEANING)
    write_json(unit_path, unit)
    meta = read(data / 'meta.json')
    meta['generatedAt'] = datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ')
    for name in ['words', 'senses', 'examples']:
        meta['counts'][name] += 1
    meta.update(wordsHash=sha(data / 'words.json'), contentHash=published_content_hash(data))
    write_json(data / 'meta.json', meta)
    catalog = [dict(wordId=row['wordId'], word=row['word'], variants=row.get('wordVariants') or []) for row in read(data / 'words.json')]
    (repo / 'src/features/direct/wordCatalog.json').write_text(json.dumps(catalog, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    evidence['packCounts'] = build_sa_pack(data, repo / 'src/features/direct')
    manifest = dict(schemaVersion=1, approvedAt='2026-10-09', approvalMessageId=APPROVAL_ID,
                    approvalScope='nostalgic 主詞、懷念往日時光義與原工作例句；不含圖片',
                    wordId=WORD_ID, senseId=SENSE_ID, exampleId=EXAMPLE_ID,
                    learningItemIds=[item['learningItemId']], originalExample=example,
                    targetMeaningZh=MEANING, originalChapterMeaningsZh=original_chapter_meaning,
                    illustrationApproval='not_approved', master=evidence)
    write_json(repo / 'src/features/vocabulary/lv3NostalgicReview.json', manifest)
    write_json(out / 'apply-result.json', evidence)
    return evidence


def apply_review(repo, master, out):
    assert out.resolve() != (repo / 'output').resolve() and out.resolve().is_relative_to((repo / 'output').resolve()), '--out must be dedicated beneath repo/output'
    direct = Path('src/features/direct')
    inputs = [Path('public/data/v1') / name for name in CONTENT_FILES + ['meta.json']]
    inputs += [direct / name for name in ['curriculum.json', 'curriculumUnit2.json']]
    inputs += [path.relative_to(repo) for path in (repo / direct).glob('curriculumLV4Unit*.json')]
    inputs += [Path('src/features/vocabulary/lv1ReviewedImages.json')]
    destinations = [repo / name for name in OUTPUT_FILES] + [out / master.name, out / 'apply-result.json']
    assert master.resolve() not in {path.resolve() for path in destinations}
    before = {path: path.read_bytes() if path.exists() else None for path in destinations}
    source_before = {repo / path: (repo / path).read_bytes() for path in inputs}
    source_before[master] = master.read_bytes()
    with tempfile.TemporaryDirectory(prefix='.nostalgic-stage-', dir=repo.parent) as directory:
        root = Path(directory).resolve()
        assert root.is_relative_to(repo.parent.resolve())
        stage = root / 'repo'
        for path in inputs:
            target = stage / path
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(repo / path, target)
        stage_out = root / 'source'
        evidence = build_outputs(stage, master, stage_out)
        assert all(path.read_bytes() == original for path, original in source_before.items()), 'Input changed; stop'
        outputs = [(stage / name, repo / name) for name in OUTPUT_FILES]
        outputs += [(stage_out / master.name, out / master.name), (stage_out / 'apply-result.json', out / 'apply-result.json')]
        replace_outputs(outputs, before)
        return evidence


def publish_master(master, out, repo=None):
    repo = repo or Path(__file__).resolve().parent.parent
    evidence = read(out / 'apply-result.json')
    audit = read(out.parent / 'scope-audit.json')
    staged = out / master.name
    assert audit['status'] == 'pass' and audit['master']['sha256'] == evidence['stagedSha256'] == sha(staged)
    assert audit['appSha256'] == {name: sha(repo / name) for name in OUTPUT_FILES}, 'Audited App outputs changed; stop'
    assert sha(master) == BASE_MASTER_HASH, 'Source changed; do not overwrite'
    backup = master.with_name(master.stem + '_before_nostalgic_20261009.xlsx')
    assert not backup.exists(), 'Backup already exists; inspect instead of overwriting'
    original = master.read_bytes()
    with backup.open('xb') as stream:
        stream.write(original)
    assert sha(backup) == BASE_MASTER_HASH
    replace_outputs([(staged, master)], {master: original})
    assert sha(master) == evidence['stagedSha256']
    return dict(published=True, backup=str(backup), sourceSha256=BASE_MASTER_HASH, masterSha256=sha(master))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--master', type=Path, required=True)
    parser.add_argument('--out', type=Path, required=True)
    parser.add_argument('--publish-master', action='store_true')
    args = parser.parse_args()
    repo = Path(__file__).resolve().parent.parent
    result = publish_master(args.master, args.out) if args.publish_master else apply_review(repo, args.master, args.out)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
