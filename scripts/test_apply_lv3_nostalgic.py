"""Exercise nostalgic staging and source publication in disposable fixtures."""
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl

import apply_lv3_nostalgic as apply

ROOT = Path(__file__).resolve().parent.parent


class NostalgicApplyTest(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name).resolve()
        self.repo = self.root / 'repo'
        direct = Path('src/features/direct')
        files = [Path('public/data/v1') / name for name in apply.CONTENT_FILES + ['meta.json', 'sa-pack.json']]
        files += [direct / name for name in ['curriculum.json', 'curriculumUnit2.json', 'wordCatalog.json']]
        files += [path.relative_to(ROOT) for path in (ROOT / direct).glob('curriculumLV4Unit*.json')]
        files += [Path('src/features/vocabulary/lv1ReviewedImages.json')]
        for path in files:
            target = self.repo / path
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(ROOT / path, target)
        for filename, key, value in [('words.json', 'wordId', apply.WORD_ID), ('senses.json', 'senseId', apply.SENSE_ID), ('examples.json', 'exampleId', apply.EXAMPLE_ID)]:
            path = self.repo / 'public/data/v1' / filename
            apply.write_json(path, [row for row in apply.read(path) if row[key] != value])
        catalog = self.repo / direct / 'wordCatalog.json'
        apply.write_json(catalog, [row for row in apply.read(catalog) if row['wordId'] != apply.WORD_ID])
        unit_path = self.repo / direct / 'curriculum.json'
        unit = apply.read(unit_path)
        item = next(row for row in unit['learningItems'] if row.get('displayWord') == apply.WORD)
        item.update(lexemeRef='LC-LV3U01-nostalgic', officialWordId=None, officialSenseId=None, chapterMeaningsZh='鄉愁的')
        apply.write_json(unit_path, unit)
        data = self.repo / 'public/data/v1'
        meta = apply.read(data / 'meta.json')
        for name in ['words', 'senses', 'examples']:
            meta['counts'][name] = len(apply.read(data / (name + '.json')))
        apply.write_json(data / 'meta.json', meta)
        apply.build_sa_pack(data, self.repo / direct)
        self.master = self.root / 'master.xlsx'
        self.master.write_bytes(b'original synthetic source')
        self.out = self.repo / 'output/source'
        self.out.mkdir(parents=True)
        (self.out / self.master.name).write_bytes(b'previous staged workbook')
        (self.out / 'apply-result.json').write_bytes(b'previous receipt')

    def stage_master(self, _source, staged):
        workbook = openpyxl.Workbook()
        words = workbook.active
        words.title = 'input_words_單字主表'
        words.append(['word_id'])
        words.append(['可空；程式自動產生'])
        words.append([apply.WORD_ID, apply.WORD, 'LV3', 'adj.', apply.MEANING])
        senses = workbook.create_sheet('input_senses_義項表')
        senses.append(['sense_id'])
        senses.append([apply.SENSE_ID, apply.WORD_ID, apply.WORD, 'adj.', apply.MEANING, None, None, apply.WORD, None, 'reviewed'])
        examples = workbook.create_sheet('input_examples_例句表')
        examples.append(['example_id'])
        examples.append(['可空；程式自動產生'])
        examples.append([apply.EXAMPLE_ID, apply.WORD, 'adj.', apply.MEANING, 'daily', apply.SENTENCE, apply.SENTENCE_ZH,
                         'Finding her old school notebook made her feel _____.', apply.WORD, 'LV3', 'reviewed'])
        workbook.save(staged)
        workbook.close()
        return dict(stagedSha256=apply.sha(staged))

    def run_apply(self):
        with patch.object(apply, 'stage_master', side_effect=self.stage_master):
            return apply.apply_review(self.repo, self.master, self.out)

    def snapshot(self):
        return {path.relative_to(self.root).as_posix(): apply.sha(path) for path in self.root.rglob('*') if path.is_file()}

    def assert_failure_unchanged(self, error):
        before = self.snapshot()
        with self.assertRaises(error):
            self.run_apply()
        self.assertEqual(self.snapshot(), before)

    def test_late_pack_failure_preserves_every_file_and_allows_retry(self):
        original = apply.build_sa_pack
        def fail_after_generation(*args):
            original(*args)
            raise RuntimeError('Late pack failure')
        with patch.object(apply, 'build_sa_pack', side_effect=fail_after_generation):
            self.assert_failure_unchanged(RuntimeError)
        self.assertEqual(self.run_apply()['packCounts']['words'], 1530)
        self.assertEqual(self.master.read_bytes(), b'original synthetic source')
        before = self.snapshot()
        with self.assertRaisesRegex(AssertionError, 'Already applied'):
            self.run_apply()
        self.assertEqual(self.snapshot(), before)

    def test_original_sentence_change_aborts_without_any_output_change(self):
        path = self.repo / 'src/features/direct/curriculum.json'
        unit = apply.read(path)
        next(row for row in unit['learningItems'] if row.get('displayWord') == apply.WORD)['originalExample']['sentenceEn'] = 'Unapproved replacement'
        apply.write_json(path, unit)
        self.assert_failure_unchanged(AssertionError)

    def test_replacement_failure_removes_new_manifest_and_restores_old_files(self):
        from apply_lv3_approved_review import os
        original = os.replace
        calls = 0
        def fail_once(source, destination):
            nonlocal calls
            calls += 1
            if calls == 9:
                raise OSError('Replacement failure after new manifest')
            return original(source, destination)
        with patch.object(os, 'replace', side_effect=fail_once):
            self.assert_failure_unchanged(OSError)

    def prepare_publish(self):
        result = self.run_apply()
        apply.write_json(self.out.parent / 'scope-audit.json', dict(status='pass', master=dict(sha256=result['stagedSha256']),
                         appSha256={name: apply.sha(self.repo / name) for name in apply.OUTPUT_FILES}))

    def test_source_replacement_failure_keeps_original_and_valid_backup(self):
        self.prepare_publish()
        from apply_lv3_approved_review import os
        with patch.object(apply, 'BASE_MASTER_HASH', apply.sha(self.master)), patch.object(os, 'replace', side_effect=OSError('Source replacement failed')):
            with self.assertRaises(OSError):
                apply.publish_master(self.master, self.out, self.repo)
        self.assertEqual(self.master.read_bytes(), b'original synthetic source')
        self.assertEqual(self.master.with_name('master_before_nostalgic_20261009.xlsx').read_bytes(), self.master.read_bytes())

    def test_successful_source_publish_retains_backup_and_refuses_repeat(self):
        self.prepare_publish()
        original = self.master.read_bytes()
        with patch.object(apply, 'BASE_MASTER_HASH', apply.sha(self.master)):
            result = apply.publish_master(self.master, self.out, self.repo)
            before = self.snapshot()
            with self.assertRaises(AssertionError):
                apply.publish_master(self.master, self.out, self.repo)
            self.assertEqual(self.snapshot(), before)
        self.assertTrue(result['published'])
        self.assertEqual(Path(result['backup']).read_bytes(), original)
        self.assertEqual(self.master.read_bytes(), (self.out / self.master.name).read_bytes())

    def test_audited_app_change_refuses_source_write(self):
        self.prepare_publish()
        (self.repo / 'public/data/v1/examples.json').write_bytes(b'changed after audit')
        before = self.snapshot()
        with patch.object(apply, 'BASE_MASTER_HASH', apply.sha(self.master)):
            with self.assertRaisesRegex(AssertionError, 'Audited App outputs changed'):
                apply.publish_master(self.master, self.out, self.repo)
        self.assertEqual(self.snapshot(), before)


if __name__ == '__main__':
    unittest.main()
